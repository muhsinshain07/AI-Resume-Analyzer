from io import BytesIO
from pathlib import Path

class ResumeParseError(Exception):
    """Raised when a resume cannot be read."""

def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as exc:
        raise ResumeParseError("Could not read the PDF resume.") from exc

def _docx(data: bytes) -> str:
    try:
        from docx import Document
        document = Document(BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ResumeParseError("Could not read the DOCX resume.") from exc

def extract_resume_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _pdf(data)
    elif suffix == ".docx":
        text = _docx(data)
    else:
        raise ResumeParseError("Only PDF and DOCX files are supported.")
    if not text:
        raise ResumeParseError("No readable text was found. Scanned PDFs require OCR, which this version does not include.")
    return text
